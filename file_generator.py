import os
from docx import Document
from docx.shared import Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

DOCS_DIR = "documents"
os.makedirs(DOCS_DIR, exist_ok=True)


def create_docx(topic, content, filename=None):
    """Создаёт Word-документ с заголовком и текстом."""
    doc = Document()

    title = doc.add_heading(topic, level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for paragraph in content.split("\n"):
        line = paragraph.strip()
        if line:
            p = doc.add_paragraph(line)
            p.paragraph_format.space_after = Pt(6)

    if not filename:
        safe = "".join(c for c in topic if c.isalnum() or c in " -_")[:30].strip()
        filename = f"{safe or 'document'}.docx"

    filepath = os.path.join(DOCS_DIR, filename)
    doc.save(filepath)
    return filepath


def create_pptx(topic, slides_data, filename=None):
    """Создаёт презентацию PowerPoint с картинками и подробным текстом."""
    import urllib.request
    import urllib.parse
    from pptx import Presentation
    from pptx.util import Inches, Pt

    prs = Presentation()
    prs.slide_width = Inches(13.33)
    prs.slide_height = Inches(7.5)

    # Титульный слайд
    slide = prs.slides.add_slide(prs.slide_layouts[0])
    slide.shapes.title.text = topic
    slide.placeholders[1].text = "Сгенерировано ИИ-помощником"

    # Слайды с контентом
    for item in slides_data:
        slide = prs.slides.add_slide(prs.slide_layouts[5])

        # Заголовок
        title_box = slide.shapes.add_textbox(
            Inches(0.5), Inches(0.3), Inches(12.3), Inches(1)
        )
        title_tf = title_box.text_frame
        title_tf.text = item.get("title", "")
        title_tf.paragraphs[0].font.size = Pt(30)
        title_tf.paragraphs[0].font.bold = True

        # Текст слева
        text_box = slide.shapes.add_textbox(
            Inches(0.5), Inches(1.5), Inches(7), Inches(5.5)
        )
        tf = text_box.text_frame
        tf.word_wrap = True

        content = item.get("content", "")
        lines = [l.strip().lstrip("-•* ").strip()
                 for l in content.split("\n") if l.strip()]

        for i, line in enumerate(lines):
            if i == 0:
                p = tf.paragraphs[0]
            else:
                p = tf.add_paragraph()
            p.text = f"• {line}"
            p.font.size = Pt(14)
            p.space_after = Pt(6)

        # Картинка справа
        try:
            query = urllib.parse.quote(item.get("title", topic))
            url = f"https://source.unsplash.com/800x600/?{query}"
            img_path = os.path.join(DOCS_DIR, f"tmp_{abs(hash(query))}.jpg")
            urllib.request.urlretrieve(url, img_path)
            slide.shapes.add_picture(
                img_path,
                Inches(7.8), Inches(1.5),
                width=Inches(5), height=Inches(5)
            )
            os.remove(img_path)
        except Exception:
            pass

    # Финальный слайд
    end_slide = prs.slides.add_slide(prs.slide_layouts[0])
    end_slide.shapes.title.text = "Спасибо за внимание!"

    if not filename:
        safe = "".join(c for c in topic if c.isalnum() or c in " -_")[:30].strip()
        filename = f"{safe or 'presentation'}.pptx"

    filepath = os.path.join(DOCS_DIR, filename)
    prs.save(filepath)
    return filepath
def create_code_file(topic, code, language="python"):
    """Создаёт файл с кодом."""
    ext_map = {
        "python": ".py",
        "javascript": ".js",
        "html": ".html",
        "css": ".css",
        "java": ".java",
        "cpp": ".cpp",
        "c": ".c",
        "sql": ".sql",
        "bash": ".sh",
    }
    ext = ext_map.get(language.lower(), ".txt")

    safe = "".join(c for c in topic if c.isalnum() or c in " -_")[:30].strip()
    filename = f"{safe or 'code'}{ext}"
    filepath = os.path.join(DOCS_DIR, filename)

    with open(filepath, "w", encoding="utf-8") as f:
        f.write(code)

    return filepath